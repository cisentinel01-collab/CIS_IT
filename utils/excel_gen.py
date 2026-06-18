import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

class ExcelGenerator:
    def export_data(self, filename, headers, data, title="Report"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.sheet_view.rightToLeft = True # Arabic RTL support in Excel

        # Title
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws['A1'] = title
        ws['A1'].font = Font(size=14, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')

        # Headers
        header_fill = PatternFill(start_color="1A2A6C", end_color="1A2A6C", fill_type="solid")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col)
            cell.value = header
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        # Data
        for row_idx, row_data in enumerate(data, 3):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.alignment = Alignment(horizontal='center')

        # Auto-adjust column width
        for column_cells in ws.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = length + 2

        wb.save(filename)

    def import_items(self, filename):
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        items = []
        # Assuming headers are: Code, Name, Category, Unit, Min Stock
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            items.append({
                "code": row[0],
                "name": row[1],
                "category": row[2],
                "unit": row[3],
                "min_stock": row[4] or 0
            })
        return items
